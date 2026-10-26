"""Load a golden set from JSON, CSV or Excel. The file extension picks the reader.

JSON:  either {"meta": {"name", "version"}, "cases": [...]} or a plain list of cases.
CSV / Excel: one case per row, first row = column names. A list cell holds several
items separated by "|", e.g.  must_include = "24 hours|full refund".
conversation_history, if used, is JSON text in one cell. counterfactual is "slot | attribute",
e.g. "{passenger} | name".
Excel needs the optional extra:  pip install "aiquality[excel]"
"""

import csv
import json
from pathlib import Path

from aiquality.errors import GoldenSetError
from aiquality.test_data.golden_case import LIST_FIELDS, GoldenCase, GoldenSet

LIST_SEPARATOR = "|"


def load_golden_set(path) -> GoldenSet:
    """Read and validate every case. Raises GoldenSetError listing all problems found."""
    path = Path(path)
    if not path.is_file():
        raise GoldenSetError(str(path), [f"file not found: {path}"])

    suffix = path.suffix.lower()
    if suffix == ".json":
        name, version, rows = _read_json(path)
    elif suffix == ".csv":
        name, version, rows = path.stem, "unversioned", _read_csv(path)
    elif suffix in (".xlsx", ".xlsm"):
        name, version, rows = path.stem, "unversioned", _read_excel(path)
    else:
        raise GoldenSetError(str(path), [f"unsupported file type '{suffix}'. Use .json, .csv or .xlsx"])

    cases, problems, seen_ids = [], [], set()
    for row_number, row in enumerate(rows, start=1):
        case_id = str(row.get("id", "")).strip() if isinstance(row, dict) else ""
        if case_id and case_id in seen_ids:
            problems.append(f"case {row_number}: duplicate id '{case_id}'")
        seen_ids.add(case_id)
        if not isinstance(row, dict):
            problems.append(f"case {row_number}: must be a set of fields, got {type(row).__name__}")
            continue
        try:
            cases.append(GoldenCase.from_dict(row))
        except ValueError as error:
            problems.append(f"case {row_number}: {error}")
    if not cases and not problems:
        problems.append("no cases found")
    if problems:
        raise GoldenSetError(str(path), problems)
    return GoldenSet(name=name, version=version, cases=cases, source=str(path))


def _read_json(path: Path):
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as error:
        raise GoldenSetError(str(path), [f"not valid JSON: {error}"])
    if isinstance(data, list):
        return path.stem, "unversioned", data
    if isinstance(data, dict) and isinstance(data.get("cases"), list):
        meta = data.get("meta", {})
        return meta.get("name", path.stem), str(meta.get("version", "unversioned")), data["cases"]
    raise GoldenSetError(str(path), ["expected a list of cases or {\"meta\": ..., \"cases\": [...]}"])


def _read_csv(path: Path) -> list[dict]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return [_clean_table_row(row) for row in csv.DictReader(handle)]


def _read_excel(path: Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise GoldenSetError(str(path), ['reading Excel needs openpyxl: pip install "aiquality[excel]"'])
    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    return [_clean_table_row(dict(zip(headers, values))) for values in rows[1:]
            if any(v not in (None, "") for v in values)]


def _clean_table_row(row: dict) -> dict:
    """Spreadsheet cells are text: trim them, drop empty ones, split list columns on '|'."""
    clean = {}
    for key, value in row.items():
        if not key:
            continue
        text = "" if value is None else str(value).strip()
        if text == "":
            continue
        if key in LIST_FIELDS:
            clean[key] = [item.strip() for item in text.split(LIST_SEPARATOR) if item.strip()]
        elif key == "counterfactual":
            slot, _, attribute = text.partition(LIST_SEPARATOR)
            clean[key] = {"slot": slot.strip(), "vary": attribute.strip()} if attribute.strip() else text
        elif key == "conversation_history":
            # Earlier turns are written as JSON in one cell: [{"role": "user", "content": "..."}]
            try:
                clean[key] = json.loads(text)
            except json.JSONDecodeError:
                clean[key] = text  # GoldenCase.from_dict reports it as invalid
        else:
            clean[key] = text
    return clean
